# -*- coding: utf-8 -*-
"""Загрузка/создание Excel с правилами разноса. Используется Diadocconnector.py"""
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
import re

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
HINT_FILL = PatternFill("solid", fgColor="FFF2CC")
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)

DATE_TOKENS = [
    ("*ГГГГ.ММ.ДД*", r"\d{4}\.\d{2}\.\d{2}", "%Y.%m.%d"),
    ("*ДД.ММ.ГГГГ*", r"\d{2}\.\d{2}\.\d{4}", "%d.%m.%Y"),
    ("*ДД*.*ММ*.*ГГГГ*", r"\d{2}\.\d{2}\.\d{4}", "%d.%m.%Y"),
    ("*ДД*-*ММ*-*ГГГГ*", r"\d{2}-\d{2}-\d{4}", "%d-%m-%Y"),
    ("*ДД-ММ-ГГГГ*", r"\d{2}-\d{2}-\d{4}", "%d-%m-%Y"),
    ("*ГГГГММДД*", r"\d{8}", "%Y%m%d"),
    ("*ДДММГГГГ*", r"\d{8}", "%d%m%Y"),
    ("*ГГММДД*", r"\d{6}", "%y%m%d"),
    ("*ДДММГГ*", r"\d{6,8}", "%d%m%y"),
]

CODE_TOKENS = [
    ("*код договора*", "contract", r"[A-Za-z0-9]+"),
    ("*Код договора*", "contract", r"[A-Za-z0-9]+"),
    ("*код портфеля*", "portfolio", r".+?"),
    ("*код*", "code", r"[A-Za-z0-9]+"),
    ("*номер*", "number", r"\d+"),
    ("*uuid*", None, r"[0-9a-fA-F\-]{36}"),
]


def is_yes(value):
    return str(value or "").strip().lower() in {"да", "yes", "1", "true", "y", "+"}


def mask_to_regex(mask):
    """Маска пользователя (*ДДММГГГГ*, *код* ...) -> regex с именованными группами."""
    raw = str(mask or "").strip()
    if not raw:
        return "", "%d%m%Y"

    work = raw
    placeholders = []

    def put(regex_part):
        key = f"@@T{len(placeholders)}@@"
        placeholders.append((key, regex_part))
        return key

    date_format = "%d%m%Y"
    date_count = 0
    named_used = set()

    for token, regex_body, fmt in DATE_TOKENS:
        while token in work:
            group_name = "date" if date_count == 0 else "date_end"
            date_count += 1
            if date_count == 1:
                date_format = fmt
            work = work.replace(token, put(f"(?P<{group_name}>{regex_body})"), 1)

    for token, gname, regex_body in CODE_TOKENS:
        while token in work:
            if gname and gname not in named_used:
                named_used.add(gname)
                part = f"(?P<{gname}>{regex_body})"
            else:
                part = f"(?:{regex_body})"
            work = work.replace(token, put(part), 1)

    # оставшиеся * — любые символы; пробел и _ в маске равнозначны
    pieces = work.split("*")
    escaped_parts = []
    for piece in pieces:
        marked = re.sub(r"[ _]+", "@@SEP@@", piece)
        escaped_parts.append(re.escape(marked).replace("@@SEP@@", r"[ _]+"))
    escaped = r".*?".join(escaped_parts)
    for key, part in placeholders:
        escaped = escaped.replace(re.escape(key), part)

    prefix = r"(?:\{Документооборот завершен\})?"
    if "Документооборот завершен" in raw:
        escaped = re.sub(
            r"\\\{Документооборот\ завершен\\\}",
            prefix,
            escaped,
            count=1,
        )
        regex = "^" + escaped
    else:
        regex = "^" + prefix + escaped

    return regex, date_format


def _style_header(ws, headers):
    for col, title in enumerate(headers, 1):
        cell = ws.cell(1, col, title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    ws.row_dimensions[1].height = 32


def _write_sheet(ws, headers, rows, widths, comments=None):
    _style_header(ws, headers)
    for r_idx, row in enumerate(rows, 2):
        for c_idx, val in enumerate(row, 1):
            cell = ws.cell(r_idx, c_idx, val)
            cell.alignment = WRAP
            cell.border = THIN
        ws.row_dimensions[r_idx].height = 36
    if comments:
        for col, text in comments.items():
            ws.cell(1, col).comment = Comment(text, "Правила")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def create_default_rules_xlsx(path):
    wb = Workbook()

    ws0 = wb.active
    ws0.title = "Инструкция"
    ws0["A1"] = "Как заполнять правила разноса"
    ws0["A1"].font = Font(bold=True, size=14)
    notes = [
        "Этот файл читает скрипт Diadocconnector.py. Меняйте строки на листах — после сохранения запустите скрипт снова.",
        "Не переименовывайте листы. Пустые строки и строки без маски игнорируются.",
        "Правила проверяются сверху вниз: срабатывает ПЕРВОЕ подходящее.",
        "",
        "Маска файла — это шаблон имени. Звёздочки — подстановки:",
        "  *ДДММГГГГ*   дата 13082026",
        "  *ГГГГММДД*   дата 20260813",
        "  *ГГММДД*     дата 260813",
        "  *ДДММГГ*     дата 130826 (или 8 цифр)",
        "  *ДД.ММ.ГГГГ* дата 13.08.2026",
        "  *ГГГГ.ММ.ДД* дата 2026.08.13",
        "  *ДД-ММ-ГГГГ* дата 13-08-2026",
        "  *код*        цифры/латиница (номер договора, код брокера)",
        "  *код портфеля*   код портфеля",
        "  *код договора*   код договора",
        "  *номер*      только цифры",
        "  *uuid*       длинный идентификатор",
        "  *            любые символы",
        "Префикс {Документооборот завершен} подставлять не обязательно — скрипт его и так игнорирует.",
        "",
        "Путь назначения. Плейсхолдеры папок:",
        "  *ГГГГ* *Месяц* (08.Август 2026) *Управляющий* *Портфель* *Брокер* *Период* *Код договора* *Дата* (13.08.2026)",
        "",
        "Знак «+» в имени: файл report+.pdf и report.pdf считаются одним и тем же, повторно не копируется.",
        "",
        "Тип источника: файлы / архивы / файлы и архивы",
        "Да/Нет в колонках пишите: да  или  нет",
        "Коды портфелей: для СПУТНИК в «Имя портфеля» достаточно SPURZ 3 — номер договора скрипт допишет сам.",
    ]
    for i, text in enumerate(notes, 3):
        ws0.cell(i, 1, text)
        ws0.cell(i, 1).alignment = WRAP
        if i in (3, 21):
            ws0.cell(i, 1).fill = HINT_FILL
        ws0.row_dimensions[i].height = 18 if text else 8
    ws0.column_dimensions["A"].width = 140

    ws_src = wb.create_sheet("Источники")
    _write_sheet(
        ws_src,
        ["Управляющий", "Папка шлюза", "Тип"],
        [
            ["УК Первая", r"\\fs-01.renlife.com\alldocs\Инвестиционный департамент\7.0 Treasury\Diadoc\diadoc_connector\Документооборот завершен\7710183778-АО УК -ПЕРВАЯ-", "файлы"],
            ["Райффайзен Капитал", r"\\fs-01.renlife.com\alldocs\Инвестиционный департамент\7.0 Treasury\Diadoc\diadoc_connector\Документооборот завершен\7702358512-ООО -УК Райффайзен-", "файлы и архивы"],
            ["ТКБ Инвестмент партнерс", r"\\fs-01.renlife.com\alldocs\Инвестиционный департамент\7.0 Treasury\Diadoc\diadoc_connector\Документооборот завершен\7825489723-ТКБ Инвестмент Партнерс (АО)", "файлы и архивы"],
            ["СПУТНИК", r"\\fs-01.renlife.com\alldocs\Инвестиционный департамент\7.0 Treasury\Diadoc\diadoc_connector\Документооборот завершен\7744000951-АО -УК -СПУТНИК - УПРАВЛЕНИЕ КАПИТАЛОМ-", "файлы и архивы"],
        ],
        [28, 130, 20],
        {3: "файлы / архивы / файлы и архивы"},
    )
    dv = DataValidation(type="list", formula1='"файлы,архивы,файлы и архивы"', allow_blank=True)
    dv.add("C2:C200")
    ws_src.add_data_validation(dv)

    ws_f = wb.create_sheet("Обычные файлы")
    _write_sheet(
        ws_f,
        ["Название", "Маска файла", "Путь назначения", "Управляющий", "Пример"],
        [
            ["Выгрузка операций", "Выгрузка операций_*ДДММГГГГ*", r"10.НАПФ - Ценные бумаги\*ГГГГ*\*Управляющий*\*Месяц*", "УК Первая", "Выгрузка операций_03062026"],
            ["I02 Первая", "I02_514833_k_d_*ГГММДД*", r"2 Отчеты брокера\*ГГГГ*\*Месяц*\*Управляющий*", "УК Первая", "I02_514833_k_d_260603"],
            ["журнал учета ДС Первая", "*ГГГГ.ММ.ДД*_27011_журнал учета ДС", r"6.Журнал учета операций\*ГГГГ*\*Месяц*\*Управляющий*", "УК Первая", "2026.06.04_27011_журнал учета ДС"],
            ["Журнал учета дс Райффайзен", "*код*_Журнал учета дс_*ДДММГГГГ*", r"6.Журнал учета операций\*ГГГГ*\*Месяц*\*Управляющий*\*Портфель*", "Райффайзен Капитал", "256_Журнал учета дс_09062026"],
            ["Общий", "*ДДММГГГГ*_Общий", r"6.Журнал учета операций\*ГГГГ*\*Месяц*\*Управляющий*", "СПУТНИК", "10072026_Общий"],
            ["MC05422", "MC05422_CCX99_EQ1_*ДДММГГ*_288200832", r"6.Журнал учета операций\*ГГГГ*\*Месяц*\*Управляющий*\Сделки НКЦ", "СПУТНИК", "MC05422_CCX99_EQ1_100726_288200832"],
            ["Podtv", "Podtv_*номер* ot_*ДД-ММ-ГГГГ*_ СК*номер*", r"6.Журнал учета операций\*ГГГГ*\*Месяц*\*Управляющий*\МНО", "СПУТНИК", "Podtv_552 ot_11-08-2026_ СК4"],
            ["Отчет о движении ДС", "Отчет о движении денежных средств *ДД.ММ.ГГГГ*_*код договора*", r"6.Журнал учета операций\*ГГГГ*\*Месяц*\*Управляющий*\Сделки НКЦ\*Код договора*\*Дата*", "СПУТНИК", "Отчет о движении денежных средств 10.07.2026_26004"],
            ["Подтверждение Клиента", "Подтверждение Клиента DU.U.27011 по МНО N *номер* *ДД.ММ.ГГГГ*", r"6.Журнал учета операций\*ГГГГ*\*Месяц*\*Управляющий*", "УК Первая", "Подтверждение Клиента DU.U.27011 по МНО N 20260804002438 04.08.2026"],
            ["ЖурналУчетаДС ТКБ", "ЖурналУчетаДС_*ДД.ММ.ГГГГ*-*ДД.ММ.ГГГГ*", r"6.Журнал учета операций\*ГГГГ*\*Месяц*\*Управляющий*", "ТКБ Инвестмент партнерс", "ЖурналУчетаДС_19.08.2026-19.08.2026"],
        ],
        [28, 70, 75, 24, 55],
        {2: "Шаблон имени. См. лист Инструкция", 4: "Пусто = правило для всех управляющих"},
    )

    ws_a = wb.create_sheet("Архивы")
    _write_sheet(
        ws_a,
        ["Название", "Маска архива", "Брокер", "Управляющий в пути", "Доставать все файлы", "Путь если все файлы", "Пример"],
        [
            ["Райффайзен НАПФ", "Брокер_напф_прочие_*ДДММГГГГ*_*код портфеля*", "", "", "нет", "", "Брокер_напф_прочие_03062026_256.zip"],
            ["ТКБ журнал ДС", "ЖурналУчетаДС_*ДД.ММ.ГГГГ*-*ДД.ММ.ГГГГ*", "", "", "нет", "", "ЖурналУчетаДС_03.07.2026-03.07.2026.zip"],
            ["ТКБ отчеты брокера", "Отчеты брокера и депозитария за *ДД.ММ.ГГГГ*", "", "", "нет", "", "Отчеты брокера и депозитария за 09.07.2026.zip"],
            ["ТКБ выгрузка НАПФ", "*код портфеля*_выгрузка НАПФ_*ДДММГГ*-*ДДММГГ*", "", "ТКБ", "да", r"10.НАПФ - Ценные бумаги\*ГГГГ*\*Управляющий*\*Месяц*\*Период*", "18-225-РЖ_выгрузка НАПФ_200726-200726.zip"],
            ["СПУТНИК ВТБ", "*ДДММГГГГ* отчеты брокера ВТБ", "Брокер ВТБ", "", "нет", "", "13082026 отчеты брокера ВТБ.7z"],
            ["СПУТНИК ГПБ", "*ДДММГГГГ* отчеты брокера ГПБ", "Брокер ГПБ", "", "нет", "", "13082026 отчеты брокера ГПБ.7z"],
            ["СПУТНИК РФБ", "*ДДММГГГГ* отчеты брокера РФБ", "Брокер РФБ", "", "нет", "", "13082026 отчеты брокера РФБ.7z"],
        ],
        [22, 60, 16, 22, 20, 70, 50],
    )

    ws_af = wb.create_sheet("Файлы внутри архивов")
    _write_sheet(
        ws_af,
        ["Название", "Маска файла", "Путь назначения", "Подставлять портфель по коду", "Без портфеля", "Добавлять брокера в путь", "Пример"],
        [
            ["Выгрузка НАПФ", "*код* Выгрузка НАПФ *ГГГГММДД*.xml", r"10.НАПФ - Ценные бумаги\*ГГГГ*\*Управляющий*\*Портфель*\*Месяц*", "нет", "нет", "нет", "256  Выгрузка НАПФ 20260603.xml"],
            ["trades", "*код*_trades_*ГГГГММДД*_*", r"2 Отчеты брокера\*ГГГГ*\*Месяц*\*Управляющий*\*Портфель*", "нет", "нет", "нет", "82748_trades_20260604_063304"],
            ["brok_rpt", "brok_rpt_*код*_*ГГГГММДД*_*_final", r"2 Отчеты брокера\*ГГГГ*\*Месяц*\*Управляющий*\*Портфель*", "да", "нет", "нет", "brok_rpt_30046_20260709_071313_final"],
            ["ЖурналУчетаДС", "ЖурналУчетаДС_*ДД.ММ.ГГГГ*-*ДД.ММ.ГГГГ*", r"6.Журнал учета операций\*ГГГГ*\*Месяц*\*Управляющий*", "нет", "да", "нет", "ЖурналУчетаДС_03.07.2026-03.07.2026"],
            ["i02 k_d", "i02_*код*_k_d_*ГГММДД*", r"2 Отчеты брокера\*ГГГГ*\*Месяц*\*Управляющий*\*Портфель*", "да", "нет", "нет", "i02_42303_k_d_260709"],
            ["i02", "i02_*код*_k_*ГГММДД*", r"2 Отчеты брокера\*ГГГГ*\*Месяц*\*Управляющий*\*Портфель*", "да", "нет", "нет", "i02_50372_k_260701"],
            ["Отчет_брокера", "Отчет_брокера_*код*_счет_SD*_ДДММГГ*_*", r"2 Отчеты брокера\*ГГГГ*\*Месяц*\*Управляющий*\*Портфель*", "да", "нет", "нет", "Отчет_брокера_88955_счет_SD88955_090726_0710040525"],
            ["Отчет ВТБ", "*код*_*ГГГГММДД*_*ГГГГММДД*_*uuid*", r"2 Отчеты брокера\*ГГГГ*\*Месяц*\*Управляющий*\*Брокер*\*Портфель*", "да", "нет", "да", "614BR6E1_20260813_20260813_e268d9fe-9424-41d4-8edf-b5a436f79846.pdf"],
        ],
        [18, 55, 70, 28, 16, 24, 55],
    )

    code_rows = []
    tkb = [
        ("30046", "18-225-РЖ_42303_30046"), ("42303", "18-225-РЖ_42303_30046"),
        ("81467", "18-226-РЖ_42307_81467_88955"), ("42307", "18-226-РЖ_42307_81467_88955"),
        ("88955", "18-226-РЖ_42307_81467_88955"), ("81766", "20-228-РЖ_97900_81766"),
        ("97900", "20-228-РЖ_97900_81766"), ("81767", "20-229-РЖ_97902_81767"),
        ("97902", "20-229-РЖ_97902_81767"), ("104011", "18-230-РЖ_104011_88956"),
        ("88956", "18-230-РЖ_104011_88956"),
    ]
    for code, folder in tkb:
        code_rows.append(["ТКБ Инвестмент партнерс", "", code, folder])
    vtb = [
        ("614BR6E1", "SPURZ 1"), ("614BR001", "SPURZ 2"), ("614BRHB1", "SPURZ 3"),
        ("614D2161", "SPURZ 4"), ("614D2241", "SPURZ 10"), ("614BW1W1", "SPURZ 11"),
        ("614BW2M1", "SPURZ 12"), ("614D22N1", "SPURZ 13"), ("614BW0X1", "SPURZ 15"),
    ]
    for code, name in vtb:
        code_rows.append(["СПУТНИК", "ВТБ", code, name])
    gpb = [
        ("50372", "SPURZ"), ("50364", "SPURZ 1"), ("50365", "SPURZ 2"), ("50366", "SPURZ 3"),
        ("50367", "SPURZ 4"), ("50371", "SPURZ 5"), ("82235", "SPURZ 10"), ("96755", "SPURZ 11"),
        ("96761", "SPURZ 12"), ("133225", "SPURZ 13"), ("133229", "SPURZ 14"), ("179447", "SPURZ 15"),
    ]
    for code, name in gpb:
        code_rows.append(["СПУТНИК", "ГПБ", code, name])
    rfb = [
        ("82550", "SPURZ"), ("82542", "SPURZ 1"), ("82438", "SPURZ 2"), ("82439", "SPURZ 3"),
        ("82544", "SPURZ 4"), ("82547", "SPURZ 5"), ("82546", "SPURZ 10"), ("82545", "SPURZ 11"),
        ("82548", "SPURZ 12"), ("82549", "SPURZ 13"), ("82543", "SPURZ 14"),
    ]
    for code, name in rfb:
        code_rows.append(["СПУТНИК", "РФБ", code, name])

    ws_c = wb.create_sheet("Коды портфелей")
    _write_sheet(
        ws_c,
        ["Управляющий", "Брокер", "Код из имени файла", "Имя портфеля"],
        code_rows,
        [28, 12, 24, 40],
        {4: "Для СПУТНИК достаточно SPURZ 3 — договор допишется автоматически"},
    )

    wb.save(path)


def _cell(ws, row, col):
    val = ws.cell(row, col).value
    if val is None:
        return ""
    return str(val).strip()


def _iter_data_rows(ws):
    for r in range(2, ws.max_row + 1):
        if any(_cell(ws, r, c) for c in range(1, ws.max_column + 1)):
            yield r


def load_rules_from_xlsx(path):
    wb = load_workbook(path, data_only=True)
    result = {
        "sources": [],
        "file_rules": [],
        "archive_rules": [],
        "archive_file_rules": [],
        "manager_mapping": {},
        "tkb_codes": {},
        "sputnik_vtb": {},
        "sputnik_gpb": {},
        "sputnik_rfb": {},
    }

    type_map = {"файлы": "files", "архивы": "archives", "файлы и архивы": "mixed", "files": "files", "archives": "archives", "mixed": "mixed"}

    ws = wb["Источники"]
    for r in _iter_data_rows(ws):
        manager = _cell(ws, r, 1)
        source_dir = _cell(ws, r, 2)
        src_type = type_map.get(_cell(ws, r, 3).lower(), "files")
        if not source_dir:
            continue
        result["sources"].append({"source_dir": source_dir, "manager": manager, "type": src_type})
        folder = source_dir.replace("/", "\\").rstrip("\\").split("\\")[-1]
        if manager:
            result["manager_mapping"][folder] = manager

    ws = wb["Обычные файлы"]
    for r in _iter_data_rows(ws):
        ident = _cell(ws, r, 1) or f"file_{r}"
        mask = _cell(ws, r, 2)
        dest = _cell(ws, r, 3)
        manager = _cell(ws, r, 4)
        if not mask or not dest:
            continue
        regex, date_fmt = mask_to_regex(mask)
        result["file_rules"].append({
            "identifier": ident,
            "mask": mask,
            "name_pattern": "",
            "date_regex": regex,
            "date_format": date_fmt,
            "destination_template": dest,
            "manager": manager,
        })

    ws = wb["Архивы"]
    for r in _iter_data_rows(ws):
        ident = _cell(ws, r, 1) or f"archive_{r}"
        mask = _cell(ws, r, 2)
        broker = _cell(ws, r, 3)
        manager_override = _cell(ws, r, 4)
        extract_all = is_yes(_cell(ws, r, 5))
        dest_all = _cell(ws, r, 6)
        if not mask:
            continue
        regex, date_fmt = mask_to_regex(mask)
        result["archive_rules"].append({
            "identifier": ident,
            "archive_pattern": "",
            "archive_regex": regex,
            "archive_date_format": date_fmt,
            "archive_date_format_8": "%d%m%Y",
            "broker": broker or None,
            "manager_override": manager_override or None,
            "extract_all": extract_all,
            "destination_template": dest_all or None,
            "kind": ident,
        })

    ws = wb["Файлы внутри архивов"]
    for r in _iter_data_rows(ws):
        ident = _cell(ws, r, 1) or f"inner_{r}"
        mask = _cell(ws, r, 2)
        dest = _cell(ws, r, 3)
        use_map = is_yes(_cell(ws, r, 4))
        no_port = is_yes(_cell(ws, r, 5))
        needs_broker = is_yes(_cell(ws, r, 6))
        if not mask or not dest:
            continue
        regex, date_fmt = mask_to_regex(mask)
        result["archive_file_rules"].append({
            "identifier": ident,
            "file_pattern": regex,
            "file_date_format": date_fmt,
            "destination_template": dest,
            "use_code_map": use_map,
            "no_portfolio": no_port,
            "needs_broker": needs_broker,
        })

    ws = wb["Коды портфелей"]
    for r in _iter_data_rows(ws):
        manager = _cell(ws, r, 1)
        broker = _cell(ws, r, 2)
        code = _cell(ws, r, 3)
        name = _cell(ws, r, 4)
        if not code:
            continue
        if "ТКБ" in manager:
            result["tkb_codes"][code] = name or code
        elif "СПУТНИК" in manager.upper() or manager.upper() == "СПУТНИК":
            if "ГПБ" in broker.upper():
                result["sputnik_gpb"][code] = name or "SPURZ"
            elif "РФБ" in broker.upper():
                result["sputnik_rfb"][code] = name or "SPURZ"
            else:
                result["sputnik_vtb"][code] = name or "SPURZ"
        else:
            result["tkb_codes"][code] = name or code

    return result
