# -*- coding: utf-8 -*-
import importlib.util
import re
from pathlib import Path

import openpyxl

ROOT = Path(r"c:\Users\Ilya.Matveev2\ProjectDocenko")
XLSX = ROOT / "Мэппинг файлов по Управляющим.xlsx"
OUT = ROOT / "coverage_report.txt"
SCRIPT = ROOT / "Diadocconnector.py"

spec = importlib.util.spec_from_file_location("d", str(SCRIPT))
d = importlib.util.module_from_spec(spec)
spec.loader.exec_module(d)

wb = openpyxl.load_workbook(str(XLSX), data_only=True)
ws = wb.active


GATEWAY_PREFIX = "{Документооборот завершен}"


def candidates_from_example(example):
    if not example:
        return []
    result = []
    # "i02_50372_k_260701 (i02_50372_k_d_260701)" -> both variants
    parts = re.split(r"[()]", str(example))
    for p in parts:
        p = p.strip()
        if not p:
            continue
        base_variants = [p]
        if not p.startswith(GATEWAY_PREFIX):
            base_variants.append(GATEWAY_PREFIX + p)
        for base in base_variants:
            result.append(base)
            if not re.search(r"\.[A-Za-z0-9]{2,5}$", base):
                for ext in (".pdf", ".xml", ".xlsx", ".xls", ".csv", ".txt"):
                    result.append(base + ext)
    # unique preserve order
    seen = set()
    uniq = []
    for item in result:
        if item not in seen:
            seen.add(item)
            uniq.append(item)
    return uniq


def archive_name_candidates(h_value, example, archive_mask):
    names = []
    h = str(h_value or "")
    if h.startswith("Архив "):
        base = h[len("Архив "):].strip()
        names.extend([base, base + ".zip", "{Документооборот завершен}" + base, "{Документооборот завершен}" + base + ".zip"])
    if example:
        ex = str(example).strip()
        names.extend([ex, ex + ".zip", "{Документооборот завершен}" + ex, "{Документооборот завершен}" + ex + ".zip"])
    # Маски вида "*ДДММГГГГ* отчеты брокера ВТБ" — подставляем дату из примера H
    if archive_mask and "отчеты брокера" in str(archive_mask):
        m = re.search(r"(\d{8})\s+отчеты брокера", h)
        if m:
            built = f"{m.group(1)} отчеты брокера " + str(archive_mask).split("отчеты брокера")[-1].strip().replace("*", "").strip()
            # better: take broker from mask
            broker = None
            for bname in ("ВТБ", "ГПБ", "РФБ"):
                if bname in str(archive_mask):
                    broker = bname
            if broker and m:
                built = f"{m.group(1)} отчеты брокера {broker}"
                names.extend([built, built + ".zip", "{Документооборот завершен}" + built + ".zip"])
        # from H like "Архив 03062026 отчеты брокера ВТБ"
        m2 = re.search(r"(\d{8})\s+отчеты брокера\s+(\S+)", h)
        if m2:
            built = f"{m2.group(1)} отчеты брокера {m2.group(2)}"
            names.extend([built, built + ".zip", "{Документооборот завершен}" + built + ".zip"])
    if archive_mask and "Отчеты брокера и депозитария" in str(archive_mask):
        m = re.search(r"(\d{2}\.\d{2}\.\d{4})", h)
        if m:
            built = f"Отчеты брокера и депозитария за {m.group(1)}"
            names.extend([built, built + ".zip", "{Документооборот завершен}" + built + ".zip"])
    if archive_mask and "ЖурналУчетаДС" in str(archive_mask) and example:
        names.extend([str(example), str(example) + ".zip"])
    # unique preserve order
    seen = set()
    uniq = []
    for n in names:
        if n not in seen:
            seen.add(n)
            uniq.append(n)
    return uniq


def match_file_rule(example):
    for c in candidates_from_example(example):
        for rule in d.FILE_RULES:
            date_obj, port, contract, found = d.extract_date_from_filename(c, rule)
            if found:
                return rule["identifier"], c, date_obj
    return None, None, None


def match_archive_rule(h_value, example, archive_mask):
    for an in archive_name_candidates(h_value, example, archive_mask):
        info = d.extract_archive_info(an)
        if info:
            return an, info
    return None, None


def match_archive_file_rule(example):
    for c in candidates_from_example(example):
        file_date, file_port, broker_code, rule = d.extract_archive_file_info(c)
        if rule:
            return rule["identifier"], c, broker_code, file_date
    return None, None, None, None


covered = []
partial = []
missing = []

for r in range(4, ws.max_row + 1):
    a = ws[f"A{r}"].value
    b = ws[f"B{r}"].value
    c = ws[f"C{r}"].value
    dd = ws[f"D{r}"].value
    e = ws[f"E{r}"].value
    f = ws[f"F{r}"].value
    g = ws[f"G{r}"].value
    h = ws[f"H{r}"].value
    i = ws[f"I{r}"].value
    if not dd and not h:
        continue

    source_tail = str(a).split("\\")[-1] if a else ""
    manager = str(b) if b else ""
    is_archive = str(h or "").startswith("Архив") or str(h or "") == "Архив"

    # H="Исходный документ" -> file rules
    # H starts with Архив -> archive + inner file
    if str(h or "").strip() == "Исходный документ":
        rid, matched, date_obj = match_file_rule(dd)
        row = {
            "row": r,
            "manager": manager,
            "portfolio": c,
            "example": dd,
            "mask": e,
            "source": source_tail,
            "h": h,
        }
        if rid:
            row["detail"] = f"FILE_RULE «{rid}» матчит «{matched}», дата={date_obj.strftime('%d.%m.%Y')}"
            covered.append(row)
        else:
            row["detail"] = "Нет подходящего FILE_RULE"
            missing.append(row)
        continue

    # Archive flow
    arch_name, arch_info = match_archive_rule(h, dd, i)
    file_rid, matched_file, code, file_date = match_archive_file_rule(dd)

    row = {
        "row": r,
        "manager": manager,
        "portfolio": c,
        "example": dd,
        "mask": e,
        "source": source_tail,
        "h": h,
        "archive_mask": i,
    }

    if arch_info and arch_info.get("extract_all"):
        row["detail"] = (
            f"ARCHIVE extract_all «{arch_name}», kind={arch_info.get('kind')}, "
            f"period={arch_info.get('period')}, дата={arch_info['date'].strftime('%d.%m.%Y')}"
        )
        covered.append(row)
        continue

    if arch_info and file_rid:
        mapped = None
        if code:
            mgr_name = d.MANAGER_MAPPING.get(source_tail) or (
                "СПУТНИК" if "СПУТНИК" in manager else
                "ТКБ Инвестмент партнерс" if "ТКБ" in manager else manager
            )
            # prefer B mapping
            if "СПУТНИК" in manager:
                mgr_name = "СПУТНИК"
            elif "ТКБ" in manager:
                mgr_name = "ТКБ Инвестмент партнерс"
            elif "Райффайзен" in manager:
                mgr_name = "Райффайзен Капитал"
            mapped = d.resolve_portfolio_folder(str(code), mgr_name, arch_info.get("broker"))
        row["detail"] = (
            f"ARCHIVE «{arch_name}» + FILE_RULE «{file_rid}» "
            f"(файл «{matched_file}», код={code}, портфель={mapped}, "
            f"брокер={arch_info.get('broker')}, дата={file_date.strftime('%d.%m.%Y')})"
        )
        covered.append(row)
        continue

    if arch_info and not file_rid:
        row["detail"] = f"Архив распознан («{arch_name}»), но пример файла не матчится ARCHIVE_FILE_RULES"
        partial.append(row)
        continue

    if (not arch_info) and file_rid:
        row["detail"] = f"Файл матчится «{file_rid}», но имя архива из H/I не распознано ARCHIVE_RULES"
        partial.append(row)
        continue

    row["detail"] = "Не распознаны ни архив, ни файл"
    missing.append(row)


def write_section(fh, title, rows):
    fh.write("\n")
    fh.write("=" * 90 + "\n")
    fh.write(f"{title} ({len(rows)})\n")
    fh.write("=" * 90 + "\n")
    if not rows:
        fh.write("(пусто)\n")
        return
    for row in rows:
        fh.write(f"\nСтрока Excel: {row['row']}\n")
        fh.write(f"  Управляющий: {row['manager']}\n")
        fh.write(f"  Портфель:    {row['portfolio']}\n")
        fh.write(f"  Пример файла:{row['example']}\n")
        fh.write(f"  Маска:       {row['mask']}\n")
        fh.write(f"  Источник H:  {row['h']}\n")
        if row.get("archive_mask"):
            fh.write(f"  Маска архива:{row['archive_mask']}\n")
        fh.write(f"  Папка шлюза: {row['source']}\n")
        fh.write(f"  Итог:        {row['detail']}\n")


with open(OUT, "w", encoding="utf-8") as fh:
    fh.write("Отчёт покрытия мэппинга правилами Diadocconnector.py\n")
    fh.write(f"Файл мэппинга: {XLSX.name}\n")
    fh.write(f"Проверка учитывает префикс шлюза: {GATEWAY_PREFIX}\n")
    fh.write(f"Всего строк данных: {len(covered) + len(partial) + len(missing)}\n")
    fh.write(f"Покрыто полностью: {len(covered)}\n")
    fh.write(f"Частично:          {len(partial)}\n")
    fh.write(f"Не покрыто:        {len(missing)}\n")
    write_section(fh, "ПОКРЫТО (скрипт умеет обработать)", covered)
    write_section(fh, "ЧАСТИЧНО (архив или файл распознан не полностью)", partial)
    write_section(fh, "НЕ ПОКРЫТО (скрипт пропустит)", missing)

print(f"OK -> {OUT}")
print(f"covered={len(covered)} partial={len(partial)} missing={len(missing)}")
